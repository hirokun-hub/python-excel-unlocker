import pytest
import boto3
import os
from moto import mock_s3
from unittest.mock import patch

# Test environment setup
@pytest.fixture(scope="session", autouse=True)
def setup_test_env():
    """Set up test environment variables"""
    os.environ.update({
        'AWS_DEFAULT_REGION': 'us-east-1',
        'AWS_ACCESS_KEY_ID': 'testing',
        'AWS_SECRET_ACCESS_KEY': 'testing',
        'AWS_SECURITY_TOKEN': 'testing',
        'AWS_SESSION_TOKEN': 'testing',
        'S3_BUCKET_NAME': 'test-bucket',
        'LOG_LEVEL': 'DEBUG'
    })

@pytest.fixture
def aws_credentials():
    """Mocked AWS Credentials for moto."""
    os.environ['AWS_ACCESS_KEY_ID'] = 'testing'
    os.environ['AWS_SECRET_ACCESS_KEY'] = 'testing'
    os.environ['AWS_SECURITY_TOKEN'] = 'testing'
    os.environ['AWS_SESSION_TOKEN'] = 'testing'
    os.environ['AWS_DEFAULT_REGION'] = 'us-east-1'

@pytest.fixture
def s3_client(aws_credentials):
    """Create a mocked S3 client"""
    with mock_s3():
        client = boto3.client('s3', region_name='us-east-1')
        # Create test bucket
        client.create_bucket(Bucket='test-bucket')
        yield client

@pytest.fixture
def sample_excel_file():
    """Create a sample Excel file for testing"""
    import io
    import openpyxl
    
    # Create a simple Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'Test Data'
    ws['B1'] = 'Value'
    ws['A2'] = 'Sample'
    ws['B2'] = 123
    
    # Save to bytes
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    return excel_buffer.getvalue()

@pytest.fixture
def password_protected_excel():
    """Create a password-protected Excel file for testing"""
    import io
    import msoffcrypto
    import openpyxl
    
    # Create a simple Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'Protected Data'
    ws['B1'] = 'Secret Value'
    
    # Save to bytes
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    # Encrypt with password
    encrypted_buffer = io.BytesIO()
    office_file = msoffcrypto.OfficeFile(excel_buffer)
    office_file.load_key(password='testpass')
    office_file.encrypt(encrypted_buffer)
    encrypted_buffer.seek(0)
    
    return encrypted_buffer.getvalue()

@pytest.fixture
def lambda_context():
    """Mock Lambda context object"""
    class MockContext:
        def __init__(self):
            self.function_name = 'test-function'
            self.function_version = '$LATEST'
            self.invoked_function_arn = 'arn:aws:lambda:us-east-1:123456789012:function:test-function'
            self.memory_limit_in_mb = 128
            self.remaining_time_in_millis = lambda: 30000
            self.log_group_name = '/aws/lambda/test-function'
            self.log_stream_name = '2023/01/01/[$LATEST]test'
            self.aws_request_id = 'test-request-id'
    
    return MockContext()

@pytest.fixture
def mock_logger():
    """Mock logger for testing"""
    with patch('src.main.logger') as mock_log:
        yield mock_log